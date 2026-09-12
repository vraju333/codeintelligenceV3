from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from baseline_models import ScenarioBaseline
from db_models import Scenario


class ExcelExportService:

    def generate_regression_report(self, report: dict, db: Session | None = None):
        workbook = Workbook()
        workbook.remove(workbook.active)

        self._build_summary_sheet(workbook.create_sheet("Summary"), report)

        if db is not None:
            scenarios = db.query(Scenario).order_by(Scenario.id).all()
            baselines = (
                db.query(ScenarioBaseline)
                .order_by(ScenarioBaseline.scenario_id, ScenarioBaseline.baseline_version)
                .all()
            )
            self._build_registry_sheet(workbook.create_sheet("Scenario Registry"), scenarios)
            self._build_active_baselines_sheet(workbook.create_sheet("Active Baselines"), scenarios, baselines)
            self._build_history_sheet(workbook.create_sheet("Baseline History"), baselines)
            self._build_endpoint_flows_sheet(workbook.create_sheet("Endpoint Flows"), baselines)

        self._build_methods_sheet(workbook.create_sheet("Changed Methods"), report)
        self._build_scenarios_sheet(workbook.create_sheet("Impact Analysis"), report)

        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return output

    def _build_summary_sheet(self, sheet, report: dict):
        sheet.append(["CodeIntelligence", "Scenario & Regression Report"])
        sheet.merge_cells(start_row=1, start_column=2, end_row=1, end_column=4)
        sheet["A1"].font = Font(bold=True, size=16)
        sheet["B1"].font = Font(bold=True, size=16)
        sheet.append([])
        sheet.append(["Report Status", report.get("status")])
        sheet.append(["Generated At", report.get("generated_at")])
        sheet.append([])
        sheet.append(["Metric", "Value"])
        self._style_header(sheet, 6)
        summary = report.get("summary", {})
        metrics = [
            ("Registered Scenarios", summary.get("registered_scenarios", 0)),
            ("Changed Java Files", summary.get("changed_java_files", 0)),
            ("Changed Classes", summary.get("changed_classes", 0)),
            ("Changed Methods", summary.get("changed_methods", 0)),
            ("Directly Affected", summary.get("directly_affected", 0)),
            ("Possibly Affected", summary.get("possibly_affected", 0)),
            ("Unaffected / No Baseline", summary.get("unaffected", 0)),
        ]
        for row in metrics:
            sheet.append(list(row))
        self._auto_size(sheet)

    def _build_registry_sheet(self, sheet, scenarios):
        headers = ["Scenario ID", "Scenario Code", "Scenario Name", "HTTP Method", "Endpoint", "Description", "Expected DB Effect", "Status"]
        sheet.append(headers)
        self._style_header(sheet, 1)
        for s in scenarios:
            sheet.append([s.id, s.scenario_code, s.scenario_name, s.http_method, s.endpoint, s.description, s.expected_db_effect, s.status])
        self._auto_size(sheet)

    def _build_active_baselines_sheet(self, sheet, scenarios, baselines):
        active = {b.scenario_id: b for b in baselines if b.is_active}
        headers = ["Scenario ID", "Scenario", "HTTP Method", "Endpoint", "Baseline Captured", "Active Version", "Flow Stored", "Captured At"]
        sheet.append(headers)
        self._style_header(sheet, 1)
        for s in scenarios:
            b = active.get(s.id)
            sheet.append([
                s.id, s.scenario_code, s.http_method, s.endpoint,
                "YES" if b else "NO",
                b.baseline_version if b else None,
                "YES" if b and b.endpoint_flow else "NO",
                b.created_at.isoformat() if b else None,
            ])
        self._auto_size(sheet)

    def _build_history_sheet(self, sheet, baselines):
        headers = ["Scenario ID", "Scenario", "Version", "Active", "HTTP Method", "Endpoint", "Flow Stored", "Created At"]
        sheet.append(headers)
        self._style_header(sheet, 1)
        for b in baselines:
            sheet.append([b.scenario_id, b.scenario_code, b.baseline_version, "YES" if b.is_active else "NO", b.http_method, b.endpoint, "YES" if b.endpoint_flow else "NO", b.created_at.isoformat()])
        self._auto_size(sheet)

    def _build_endpoint_flows_sheet(self, sheet, baselines):
        headers = ["Scenario", "Baseline Version", "Active", "HTTP Method", "Endpoint", "Stored Flow"]
        sheet.append(headers)
        self._style_header(sheet, 1)
        for b in baselines:
            flow = b.endpoint_flow or {}
            simplified = flow.get("simplified_flow") if isinstance(flow, dict) else None
            if simplified:
                flow_text = " -> ".join(str(x) for x in simplified)
            else:
                flow_text = json.dumps(flow, ensure_ascii=False) if flow else ""
            sheet.append([b.scenario_code, b.baseline_version, "YES" if b.is_active else "NO", b.http_method, b.endpoint, flow_text])
        self._auto_size(sheet)

    def _build_methods_sheet(self, sheet, report: dict):
        sheet.append(["Class", "Method", "Changed Lines"])
        self._style_header(sheet, 1)
        for method in report.get("changed_methods", []):
            lines = method.get("changed_lines", [])
            sheet.append([method.get("class_name"), method.get("method_name"), ", ".join(map(str, lines))])
        self._auto_size(sheet)

    def _build_scenarios_sheet(self, sheet, report: dict):
        headers = ["Scenario ID", "Scenario", "HTTP Method", "Endpoint", "Baseline Version", "Impact Status", "Matched Classes", "Matched Methods", "Reason"]
        sheet.append(headers)
        self._style_header(sheet, 1)
        for scenario in report.get("scenarios", []):
            sheet.append([
                scenario.get("scenario_id"), scenario.get("scenario_code"), scenario.get("http_method"), scenario.get("endpoint"),
                scenario.get("baseline_version"), scenario.get("impact_status"),
                ", ".join(scenario.get("matched_classes", [])),
                ", ".join(scenario.get("matched_methods", [])),
                scenario.get("reason"),
            ])
        self._auto_size(sheet)

    def _style_header(self, sheet, row_number: int):
        fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        for cell in sheet[row_number]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    def _auto_size(self, sheet):
        for column_cells in sheet.columns:
            length = 0
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                length = max(length, min(len(value), 70))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = max(12, min(length + 2, 55))
        sheet.freeze_panes = "A2"
